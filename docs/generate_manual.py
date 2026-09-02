"""Speaky OPIc App User Manual Generator"""
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

# ============================================================
# Styles
# ============================================================
style = doc.styles['Normal']
style.font.name = 'Malgun Gothic'
style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.3

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'Malgun Gothic'
    hs.font.color.rgb = RGBColor(0x21, 0x21, 0x21)

doc.styles['Heading 1'].font.size = Pt(20)
doc.styles['Heading 2'].font.size = Pt(15)
doc.styles['Heading 3'].font.size = Pt(12)

def add_table(headers, rows):
    """Add a styled table"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Light Grid Accent 1'
    # Header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
    # Rows
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()

def section_break():
    doc.add_page_break()

# ============================================================
# Cover
# ============================================================
for _ in range(6):
    doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('Speaky')
run.font.size = Pt(36)
run.bold = True
run.font.color.rgb = RGBColor(0xD4, 0x70, 0x7F)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('OPIc 학습 플랫폼 사용 설명서')
run.font.size = Pt(18)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()
ver = doc.add_paragraph()
ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = ver.add_run('Version 1.1  |  2026년 9월')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

section_break()

# ============================================================
# TOC placeholder
# ============================================================
doc.add_heading('목차', level=1)
toc_items = [
    '1. 앱 소개',
    '2. 시작하기 (회원가입 / 로그인)',
    '3. 학생 화면',
    '   3.1 홈 탭',
    '   3.2 모의고사 탭',
    '   3.3 연습 기록 탭',
    '   3.4 설정 탭',
    '   3.5 초대 코드 입력',
    '   3.6 내 토픽 설정',
    '   3.7 OPIc 핵심 표현',
    '   3.8 메시지',
    '   3.9 모의고사 진행',
    '   3.10 스크립트 연습',
    '   3.11 토픽 상세',
    '4. 강사 화면',
    '   4.1 홈 탭 (학생 목록)',
    '   4.2 반 관리 탭',
    '   4.3 초대 탭',
    '   4.4 모의고사 탭',
    '   4.5 설정 탭',
    '   4.6 학생 관리',
    '   4.7 스크립트 관리',
    '   4.8 반 관리 (상세)',
    '   4.9 메시지',
    '   4.10 OPIc 핵심 표현 관리',
    '   4.11 학원 관리 (원장 전용)',
    '5. 관리자 화면',
    '   5.1 대시보드',
    '   5.2 사용자 관리',
    '   5.3 학원 관리',
    '   5.4 랜딩 페이지 CMS',
    '   5.5 결제 관리',
    '   5.6 설정',
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    for r in p.runs:
        r.font.size = Pt(10)

section_break()

# ============================================================
# 1. App Introduction
# ============================================================
doc.add_heading('1. 앱 소개', level=1)
doc.add_paragraph(
    'Speaky는 OPIc(Oral Proficiency Interview - Computer) 시험 대비를 위한 학습 플랫폼입니다. '
    '학원(조직) 단위로 강사와 학생을 관리하며, AI 기반 피드백과 실전 모의고사를 제공합니다.'
)

doc.add_heading('주요 기능', level=3)
features = [
    '강사-학생 연계: 강사가 학생별 맞춤 스크립트 작성 및 피드백 제공',
    '실전 연습: Ava 음성(TTS)으로 질문을 듣고 녹음하여 연습',
    'AI 피드백: 스크립트 vs 실제 답변 비교 분석 (Claude AI)',
    '모의고사: 실제 OPIc CBT와 유사한 시뮬레이션 (자동 재생/녹음)',
    'OPIc 핵심 표현: 상황별 필수 표현 학습 라이브러리',
    '학습 기록: 연습 이력, 진도 추적, 통계 대시보드',
    '메시징: 강사-학생 간 메시지 발송/수신',
]
for f in features:
    doc.add_paragraph(f, style='List Bullet')

doc.add_heading('사용자 역할', level=3)
add_table(
    ['역할', '설명', '주요 권한'],
    [
        ['학생 (Student)', '학습자', '연습, 녹음, 모의고사, 기록 조회'],
        ['강사 (Teacher)', '학원 강사', '스크립트 작성, 피드백, 학생/반 관리, 초대'],
        ['원장 (Owner)', '학원 대표', '강사 관리, 구독/결제, 학원 정보 관리'],
        ['슈퍼 관리자', '플랫폼 관리자', '전체 학원/사용자/결제/CMS 관리'],
    ]
)

section_break()

# ============================================================
# 2. Getting Started
# ============================================================
doc.add_heading('2. 시작하기', level=1)

doc.add_heading('2.1 회원가입', level=2)
doc.add_paragraph('1. 앱을 실행하고 "회원가입" 버튼을 탭합니다.')
doc.add_paragraph('2. 이름, 이메일, 비밀번호(6자 이상)를 입력합니다.')
doc.add_paragraph('3. 가입 완료 후 이메일로 발송된 8자리 인증 코드를 입력합니다.')
doc.add_paragraph('4. 인증 완료 시 자동으로 로그인됩니다.')
p = doc.add_paragraph()
run = p.add_run('참고: ')
run.bold = True
p.add_run('인증 코드 입력 시 5회 실패하면 3분간 잠금됩니다. 코드 재전송도 가능합니다.')

doc.add_heading('2.2 로그인', level=2)
doc.add_paragraph('가입한 이메일과 비밀번호로 로그인합니다.')
doc.add_paragraph('비밀번호를 잊은 경우 "비밀번호 찾기"를 통해 재설정 링크를 이메일로 받을 수 있습니다.')

doc.add_heading('2.3 학원 연결 (학생)', level=2)
doc.add_paragraph('학생은 로그인 후 강사로부터 받은 6자리 초대 코드를 입력하여 학원에 소속됩니다.')
doc.add_paragraph('웹 초대 링크(speaky.co.kr/join/CODE)로도 가입 가능합니다.')

section_break()

# ============================================================
# 3. Student Screens
# ============================================================
doc.add_heading('3. 학생 화면', level=1)
doc.add_paragraph('학생은 하단 4개 탭(홈, 모의고사, 기록, 설정)과 추가 하위 화면으로 구성됩니다.')

# 3.1 Home
doc.add_heading('3.1 홈 탭', level=2)
doc.add_paragraph(
    '학생 대시보드로, 학원 연결 상태에 따라 두 가지 화면이 표시됩니다.'
)
doc.add_heading('학원 미연결 시', level=3)
doc.add_paragraph('초대 코드 입력 버튼이 표시됩니다. 코드를 입력하여 학원에 소속하세요.')
doc.add_heading('학원 연결 후', level=3)
items = [
    ('담당 강사 카드', '연결된 강사 이름과 소속 학원 정보'),
    ('일일 목표 & 스트릭', '오늘 연습 횟수 / 목표 대비 진행률, 연속 학습 일수'),
    ('통계 스트립', '전체 연습 횟수, 평균 점수, 평균 재현율 등 핵심 통계'),
    ('약점 토픽 추천', 'AI가 분석한 약한 토픽 영역 표시'),
    ('OPIc 핵심 표현', '상황별 필수 표현 학습 페이지 바로가기'),
    ('내 토픽 목록', '강사가 배정한 토픽 카드 목록 (진도율 표시)'),
]
for title, desc in items:
    p = doc.add_paragraph()
    run = p.add_run(f'{title}: ')
    run.bold = True
    p.add_run(desc)
doc.add_paragraph(
    '토픽 목록에서 [스크립트 보관] 표시가 붙은 카드는 현재 배정에서는 빠졌지만 '
    '강사가 작성한 스크립트가 남아 있는 토픽입니다. 배정된 토픽 다음에 표시되며, '
    '탭하면 기존 스크립트를 그대로 볼 수 있습니다.'
)

# 3.2 Exam Tab
doc.add_heading('3.2 모의고사 탭', level=2)
doc.add_paragraph('모의고사 허브 화면으로 3가지 시험 유형을 선택할 수 있습니다.')
add_table(
    ['시험 유형', '설명', '문항 수'],
    [
        ['모의고사', 'OPIc 실전과 동일한 서베이 기반 시험', '12~15문항'],
        ['롤플레이 콤보', '특정 시나리오 기반 콤보 롤플레이', '3문항'],
        ['레벨 테스트', '현재 실력 진단용 레벨 테스트', '5문항'],
    ]
)
doc.add_paragraph('하단에는 최근 시험 결과 미리보기가 표시됩니다.')

# 3.3 History
doc.add_heading('3.3 연습 기록 탭', level=2)
doc.add_paragraph(
    '스크립트 기반 연습 이력을 조회합니다. 각 연습 카드에는 점수, 재현율(%), '
    '녹음 시간이 표시되며, 탭하면 상세 피드백을 확인할 수 있습니다.'
)

# 3.4 Settings
doc.add_heading('3.4 설정 탭', level=2)
items = [
    ('계정 정보', '이름, 이메일, 역할, 소속 학원 표시'),
    ('테마', '시스템 / 라이트 / 다크 모드 선택'),
    ('법적 고지', '이용약관, 개인정보처리방침'),
    ('로그아웃', '계정 로그아웃'),
]
for title, desc in items:
    p = doc.add_paragraph()
    run = p.add_run(f'{title}: ')
    run.bold = True
    p.add_run(desc)

# 3.5 Connect
doc.add_heading('3.5 초대 코드 입력', level=2)
doc.add_paragraph(
    '강사로부터 받은 6자리 초대 코드를 입력하여 학원에 소속됩니다. '
    '웹 초대 링크를 통해 접속한 경우 코드가 자동으로 입력됩니다. '
    '코드 사용 시 해당 반에 자동 배정됩니다.'
)

# 3.6 Topics
doc.add_heading('3.6 내 토픽 설정', level=2)
doc.add_paragraph('OPIc Background Survey 구조에 따라 토픽을 선택합니다.')
doc.add_heading('프로필 질문 (Q1~Q3)', level=3)
doc.add_paragraph('직업, 학생 여부, 거주지 등 기본 프로필을 설정합니다.')
doc.add_heading('활동 토픽 (Q4~Q7)', level=3)
doc.add_paragraph(
    '여가, 취미, 운동, 휴가 등 4개 그룹에서 관심 토픽을 선택합니다. '
    '총 12개 이상 선택해야 저장됩니다. 자기소개, 집/거주, 이웃/동네는 자동 배정됩니다.'
)
doc.add_paragraph(
    '선택을 해제한 토픽에 강사가 작성한 스크립트가 있으면, 저장 전에 해당 토픽과 '
    '스크립트 개수를 알려주는 확인 창이 표시됩니다. 해제해도 스크립트가 사라지지는 '
    '않으며, 홈 화면 토픽 목록에 [스크립트 보관] 표시와 함께 계속 남습니다.'
)

# 3.7 Expressions
doc.add_heading('3.7 OPIc 핵심 표현', level=2)
doc.add_paragraph(
    '상황별로 정리된 OPIc 필수 표현을 학습할 수 있습니다.'
)
add_table(
    ['카테고리', '설명'],
    [
        ['도입 표현', '답변 시작 시 자연스러운 오프닝'],
        ['묘사/설명', '장소, 사물, 사람을 묘사하는 표현'],
        ['루틴/습관', '일상적인 활동을 설명하는 표현'],
        ['경험/과거', '과거 경험을 이야기하는 표현'],
        ['비교/변화', '과거와 현재를 비교하는 표현'],
        ['롤플레이', '상황극에서 사용하는 표현'],
        ['연결/전환', '답변 흐름을 이어주는 전환 표현'],
        ['필러/시간벌기', '생각할 시간을 버는 자연스러운 표현'],
        ['감정/의견', '감정과 의견을 표현하는 문장'],
        ['마무리', '답변을 자연스럽게 마무리하는 표현'],
        ['고급 수식어', 'AL 등급을 위한 고급 표현'],
    ]
)
doc.add_paragraph('각 표현에는 영어 표현, 한국어 번역, 예문, 학습 팁이 포함되며, 기초/중급/고급 레벨로 필터링 가능합니다.')

# 3.8 Messages
doc.add_heading('3.8 메시지', level=2)
doc.add_paragraph(
    '강사가 보낸 메시지를 확인합니다. 미읽은 메시지는 파란 점으로 표시되며, '
    '탭하면 자동으로 읽음 처리됩니다. 홈 화면 헤더의 메일 아이콘에 미읽은 수가 뱃지로 표시됩니다.'
)

# 3.9 Mock Exam Flow
doc.add_heading('3.9 모의고사 진행', level=2)
doc.add_paragraph('모의고사는 아래 단계로 진행됩니다:')

doc.add_heading('Step 1: 토픽 선택 (서베이)', level=3)
doc.add_paragraph('활동 토픽(Q4~Q7)에서 시험에 출제될 토픽을 선택합니다. 서베이 전략 가이드도 참고 가능합니다.')

doc.add_heading('Step 2: 난이도 자기평가', level=3)
doc.add_paragraph('레벨 1~5 중 선택합니다. 레벨에 따라 문항 수와 난이도가 조절됩니다.')

doc.add_heading('Step 3: 오리엔테이션', level=3)
doc.add_paragraph(
    '시험 규칙과 구성이 안내됩니다. '
    'Q1~Q15 문항 구성(자기소개/서베이 토픽/돌발 질문/롤플레이) 다이어그램이 표시되며, '
    '[시험 시작]을 누르면 3초 카운트다운 후 시험이 시작됩니다.'
)
doc.add_paragraph(
    '처음 시험을 보는 경우 [시험 시작]을 누를 때 음성 녹음 동의 창이 표시됩니다. '
    '동의하면 곧바로 카운트다운이 이어집니다. 시험이 시작된 뒤에는 동의 창이 뜨지 않습니다.'
)

doc.add_heading('Step 4: 시험 진행', level=3)
items = [
    '질문이 자동으로 TTS 음성 재생됩니다',
    '음성 재생 완료 후 3초 준비 카운트다운이 표시됩니다',
    '카운트다운 종료 시 자동으로 녹음이 시작됩니다',
    '녹음 완료 후 다음 문항으로 자동 진행됩니다',
    '문항별 시간 제한이 있으며, 상단에 진행률이 표시됩니다',
    '네트워크 문제 등으로 음성이 재생되지 않아도 시험은 중단되지 않습니다. '
    '화면의 질문 글을 읽고 답변하면 됩니다',
]
for item in items:
    doc.add_paragraph(item, style='List Bullet')

doc.add_heading('Step 5: 결과 처리', level=3)
doc.add_paragraph(
    '시험 종료 후 AI 평가가 진행됩니다. 업로드 -> STT(음성인식) -> AI 평가 3단계 진행률이 표시됩니다.'
)

doc.add_heading('Step 6: 결과 확인', level=3)
items = [
    'OPIc 예상 등급 (NL ~ AL)',
    'ACTFL 4차원 분석: 기능/수행, 맥락/내용, 정확성, 텍스트 유형',
    'AI 종합 평가: 강점, 개선사항, 학습 추천',
    '문항별 상세 평가 (확장 가능)',
]
for item in items:
    doc.add_paragraph(item, style='List Bullet')

# 3.10 Script Practice
doc.add_heading('3.10 스크립트 연습', level=2)
doc.add_paragraph('토픽의 질문별로 강사가 작성한 스크립트를 바탕으로 연습합니다.')

doc.add_heading('연습 유형', level=3)
add_table(
    ['유형', '설명'],
    [
        ['셰도잉', 'TTS 음성을 들으며 스크립트를 따라 읽는 연습'],
        ['한->영 연습', '한국어 번역을 보고 영어로 말하는 연습'],
        ['실전 연습', 'TTS 질문을 듣고 스크립트 없이 답변 녹음, AI 피드백 제공'],
    ]
)
doc.add_paragraph(
    '실전 연습 후에는 AI가 스크립트와 실제 답변을 비교 분석하여 점수, 재현율, 문법/발음/자연스러움 피드백을 제공합니다.'
)

# 3.11 Topic Detail
doc.add_heading('3.11 토픽 상세', level=2)
doc.add_paragraph(
    '선택한 토픽의 모든 질문(묘사/루틴/경험/비교)을 조회합니다. '
    '각 질문에 대한 스크립트 작성 여부가 표시되며, 탭하면 스크립트 상세로 이동합니다.'
)

section_break()

# ============================================================
# 4. Teacher Screens
# ============================================================
doc.add_heading('4. 강사 화면', level=1)
doc.add_paragraph('강사는 하단 5개 탭(홈, 반 관리, 초대, 모의고사, 설정)과 다양한 하위 화면으로 구성됩니다.')

# 4.1 Home
doc.add_heading('4.1 홈 탭 (학생 목록)', level=2)
doc.add_paragraph(
    '연결된 학생 목록을 표시합니다. 각 학생 카드에는 연습 횟수, 스크립트 수, 피드백 현황이 요약됩니다. '
    '주의가 필요한 학생(미활동, 피드백 미제공)은 상단에 별도로 표시됩니다. '
    '학생 카드를 탭하면 학생 상세 화면으로 이동합니다.'
)

# 4.2 Classes
doc.add_heading('4.2 반 관리 탭', level=2)
doc.add_paragraph(
    '반(클래스) 목록을 관리합니다. 반 생성, 반원 추가/제거, 반 정보 수정이 가능합니다.'
)

# 4.3 Invite
doc.add_heading('4.3 초대 탭', level=2)
doc.add_paragraph('학생을 초대하기 위한 코드를 생성하고 관리합니다.')
items = [
    ('반 선택', '초대할 반을 선택합니다 (선택사항)'),
    ('사용 횟수', '코드 사용 가능 횟수를 설정합니다 (1회 ~ 무제한)'),
    ('공유 방법', '코드 복사, 링크 공유, QR 코드 3가지 방법 제공'),
    ('사용 현황', '각 코드의 사용 횟수와 상태(활성/만료) 확인'),
]
for title, desc in items:
    p = doc.add_paragraph()
    run = p.add_run(f'{title}: ')
    run.bold = True
    p.add_run(desc)

# 4.4 Exam
doc.add_heading('4.4 모의고사 탭', level=2)
doc.add_paragraph('강사 본인도 모의고사를 응시할 수 있습니다. 학생과 동일한 모의고사 화면이 제공됩니다.')

# 4.5 Settings
doc.add_heading('4.5 설정 탭', level=2)
items = [
    ('학원 관리 (원장 전용)', '학원 정보, 강사 관리, 구독 정보'),
    ('학습 콘텐츠', 'OPIc 핵심 표현 관리'),
    ('알림 (웹)', '미읽은 알림 확인'),
    ('계정/테마', '이름, 이메일, 다크/라이트 테마'),
    ('로그아웃', '계정 로그아웃'),
]
for title, desc in items:
    p = doc.add_paragraph()
    run = p.add_run(f'{title}: ')
    run.bold = True
    p.add_run(desc)

# 4.6 Student Management
doc.add_heading('4.6 학생 관리', level=2)

doc.add_heading('학생 상세', level=3)
doc.add_paragraph(
    '특정 학생의 종합 정보를 4개 탭(토픽/스크립트/연습/시험)으로 확인합니다. '
    '토픽별 진도, 스크립트 목록, 연습 이력, 시험 결과를 각각 조회할 수 있습니다.'
)

doc.add_heading('토픽 배정', level=3)
doc.add_paragraph(
    '학생에게 OPIc 서베이 프로필(Q1~Q3)과 활동 토픽(Q4~Q7)을 배정합니다. '
    '자동 배정 토픽(자기소개, 집/거주, 이웃/동네)은 별도 표시됩니다.'
)
doc.add_paragraph(
    '이미 스크립트를 작성한 토픽의 배정을 해제하면, 저장 전에 토픽 이름과 스크립트 '
    '개수를 알려주는 확인 창이 표시됩니다. 해제해도 스크립트는 삭제되지 않고 학생 '
    '화면에 [스크립트 보관]으로 남아 학생이 계속 볼 수 있습니다.'
)

# 4.7 Script
doc.add_heading('4.7 스크립트 관리', level=2)
doc.add_paragraph('학생별 맞춤 스크립트를 작성합니다.')
doc.add_paragraph('1. 토픽 선택: 학생에게 배정된 토픽 중 선택')
doc.add_paragraph('2. 질문 선택: 해당 토픽의 질문(묘사/루틴/경험/비교) 중 선택')
doc.add_paragraph('3. 스크립트 작성: 영어 스크립트 본문과 코멘트 입력')
doc.add_paragraph('4. 작성된 스크립트는 학생에게 자동으로 공유되며, 학생이 연습할 수 있습니다.')
p = doc.add_paragraph()
run = p.add_run('참고: ')
run.bold = True
p.add_run('스크립트 쿼터는 구독 플랜에 따라 제한됩니다.')

# 4.8 Class Detail
doc.add_heading('4.8 반 관리 (상세)', level=2)
items = [
    ('반 생성', '반 이름(필수, 최대 50자)과 설명을 입력하여 생성'),
    ('반 상세', '반원 목록 확인, 멤버 추가/제거, 반 정보 수정'),
    ('멤버 추가', '아직 이 반에 속하지 않은 학생을 다중 선택하여 추가'),
]
for title, desc in items:
    p = doc.add_paragraph()
    run = p.add_run(f'{title}: ')
    run.bold = True
    p.add_run(desc)

# 4.9 Messages
doc.add_heading('4.9 메시지', level=2)
doc.add_paragraph('학생에게 메시지를 발송하고 이력을 관리합니다.')
doc.add_heading('메시지 발송', level=3)
doc.add_paragraph(
    '대상(반 전체 또는 개별 학생)을 선택하고 제목과 본문을 입력하여 발송합니다. '
    '반 전체 발송 시 해당 반의 모든 학생에게 개별 전달됩니다.'
)
doc.add_heading('발송 이력', level=3)
doc.add_paragraph('보낸 메시지 목록에서 대상, 읽음률, 발송 시간을 확인합니다.')

# 4.10 Expressions
doc.add_heading('4.10 OPIc 핵심 표현 관리', level=2)
doc.add_paragraph('설정 > 학습 콘텐츠 > OPIc 핵심 표현 관리에서 접근합니다.')

doc.add_heading('기본 데이터 복제', level=3)
doc.add_paragraph(
    '처음 접근 시 "기본 데이터 복제" 버튼이 표시됩니다. '
    '복제하면 기본 제공 표현이 학원 전용 데이터로 복사되어 자유롭게 수정할 수 있습니다.'
)

doc.add_heading('표현 관리 (CRUD)', level=3)
items = [
    '카테고리별 탭에서 해당 카테고리의 표현 목록 확인',
    '"추가" 버튼으로 새 표현 생성 (영어 표현, 한국어 번역, 예문, 팁, 난이도)',
    '연필 아이콘으로 기존 표현 수정',
    '휴지통 아이콘으로 표현 삭제',
    '카테고리 탭 길게 누르면 카테고리 삭제',
]
for item in items:
    doc.add_paragraph(item, style='List Bullet')

# 4.11 Academy Management
doc.add_heading('4.11 학원 관리 (원장 전용)', level=2)
doc.add_paragraph('원장(Owner) 역할만 접근 가능한 학원 관리 메뉴입니다.')

doc.add_heading('학원 정보', level=3)
doc.add_paragraph('학원명을 수정할 수 있습니다.')

doc.add_heading('강사 관리', level=3)
doc.add_paragraph('소속 강사 목록 조회, 강사 제거, 역할 변경(원장/강사)이 가능합니다.')

doc.add_heading('구독 정보', level=3)
items = [
    ('구독 상태', '현재 플랜, 다음 결제일, 사용량(학생/스크립트 쿼터) 확인'),
    ('플랜 변경', 'Free / Solo / Pro / Academy 플랜 변경 (업그레이드: 즉시 일할 결제, 다운그레이드: 다음 갱신 시 적용)'),
    ('결제 수단', '카드 정보 변경'),
    ('결제 이력', '월별 결제 내역 및 영수증 확인'),
    ('구독 취소', '3단계 리텐션 플로우 (사유 선택 -> 제안 -> 최종 확인)'),
]
for title, desc in items:
    p = doc.add_paragraph()
    run = p.add_run(f'{title}: ')
    run.bold = True
    p.add_run(desc)

doc.add_heading('구독 플랜 비교', level=3)
add_table(
    ['플랜', '월 요금', '학생 수', '스크립트', 'AI 피드백', 'TTS'],
    [
        ['Free', '무료', '3명', '5개', '월 5회', '제한적'],
        ['Solo', '49,900원', '15명', '50개', '무제한', '무제한'],
        ['Pro', '99,900원', '50명', '200개', '무제한', '무제한'],
        ['Academy', '299,000원', '무제한', '무제한', '무제한', '무제한'],
    ]
)
doc.add_paragraph('연간 결제 시 25% 할인이 적용됩니다. 신규 가입 시 30일 무료 체험(Solo 플랜)이 제공됩니다.')

section_break()

# ============================================================
# 5. Admin Screens
# ============================================================
doc.add_heading('5. 관리자 화면', level=1)
doc.add_paragraph('슈퍼 관리자(super_admin)만 접근 가능한 플랫폼 전체 관리 화면입니다. 6개 탭으로 구성됩니다.')

# 5.1 Dashboard
doc.add_heading('5.1 대시보드', level=2)
doc.add_paragraph('플랫폼 전체 KPI를 한눈에 확인합니다.')
add_table(
    ['지표', '설명'],
    [
        ['전체 사용자', '가입자 수 (강사/학생 분류)'],
        ['활성 사용자', '최근 7일/30일 활동 사용자'],
        ['등록 학원', '가입된 학원(조직) 수'],
        ['콘텐츠', '작성된 스크립트 수, 완료된 연습 수'],
        ['MRR / ARR', '월간/연간 반복 수익'],
        ['이탈률', '구독 취소율'],
        ['플랜 분포', '플랜별 구독자 비율'],
    ]
)

# 5.2 Users
doc.add_heading('5.2 사용자 관리', level=2)
doc.add_paragraph(
    '전체 사용자를 검색하고 역할별(슈퍼관리자/원장/강사/학생)로 필터링합니다. '
    '사용자 상세에서 역할 변경, 소속 조직 확인, 감사 로그 조회가 가능합니다.'
)

# 5.3 Academies
doc.add_heading('5.3 학원 관리', level=2)
items = [
    ('원장 초대', '새 학원을 개설할 원장 초대 코드 생성 및 관리'),
    ('학원 목록', '등록된 학원 조회 (강사/학생 수 표시)'),
    ('학원 상세', '학원 정보 수정, 멤버 관리, 구독 상태 확인, 학원 삭제'),
]
for title, desc in items:
    p = doc.add_paragraph()
    run = p.add_run(f'{title}: ')
    run.bold = True
    p.add_run(desc)

# 5.4 Landing CMS
doc.add_heading('5.4 랜딩 페이지 CMS', level=2)
doc.add_paragraph('공개 랜딩 페이지의 콘텐츠를 관리합니다.')
add_table(
    ['섹션', '설명'],
    [
        ['Hero', '메인 배너 제목/부제목/CTA 버튼'],
        ['Features', '주요 기능 소개 카드'],
        ['Pricing', '요금제 비교표'],
        ['Stats', '플랫폼 통계 수치'],
        ['Steps', '이용 절차 안내'],
        ['FAQ', '자주 묻는 질문'],
        ['Video', '소개 영상'],
        ['Roadmap', '개발 로드맵'],
    ]
)
doc.add_paragraph('각 섹션의 아이템을 추가/수정/삭제하고 순서를 변경할 수 있습니다.')

# 5.5 Billing
doc.add_heading('5.5 결제 관리', level=2)
items = [
    ('구독 통계', 'MRR, ARR, 이탈률, 구독자 수 종합 대시보드'),
    ('플랜 관리', '각 플랜의 가격, 학생/스크립트 제한, 기능 설정 수정'),
    ('결제 이력', '최근 결제 내역 조회, 환불 처리'),
]
for title, desc in items:
    p = doc.add_paragraph()
    run = p.add_run(f'{title}: ')
    run.bold = True
    p.add_run(desc)

# 5.6 Settings
doc.add_heading('5.6 설정', level=2)
doc.add_paragraph(
    '테마 설정, 로그아웃, 감사 로그(관리자 활동 기록) 조회가 가능합니다.'
)

section_break()

# ============================================================
# Appendix
# ============================================================
doc.add_heading('부록', level=1)

doc.add_heading('A. 지원 OPIc 등급', level=2)
add_table(
    ['등급', '설명'],
    [
        ['NL (Novice Low)', '단어 수준의 제한된 의사소통'],
        ['NM (Novice Mid)', '암기된 표현으로 최소한의 의사소통'],
        ['NH (Novice High)', '간단한 문장으로 기초적 의사소통'],
        ['IL (Intermediate Low)', '기본 문장으로 일상적 주제 대화 가능'],
        ['IM (Intermediate Mid)', '다양한 문장으로 익숙한 주제 대화 가능'],
        ['IH (Intermediate High)', '문단 수준으로 복잡한 주제 설명 가능'],
        ['AL (Advanced Low)', '자연스럽고 유창한 문단 수준 대화'],
    ]
)

doc.add_heading('B. 질문 유형', level=2)
add_table(
    ['유형', '한글명', '설명'],
    [
        ['Describe', '묘사/설명', '장소, 사물, 사람을 묘사'],
        ['Routine', '루틴/습관', '일상적인 활동 설명'],
        ['Experience', '경험/과거', '과거 경험 이야기'],
        ['Comparison', '비교/변화', '과거와 현재 비교'],
        ['Roleplay', '롤플레이', '특정 상황에서 역할극'],
        ['Advanced', '고급 질문', '심화 난이도 질문'],
    ]
)

doc.add_heading('C. 단축키 / 제스처', level=2)
add_table(
    ['동작', '설명'],
    [
        ['아래로 당기기', '화면 새로고침 (Pull-to-refresh)'],
        ['카드 탭', '상세 화면으로 이동'],
        ['길게 누르기', '삭제/추가 옵션 (강사 카테고리 탭 등)'],
        ['좌우 스와이프', '카테고리/필터 탭 스크롤'],
    ]
)

# Save
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'Speaky_사용설명서.docx')
doc.save(output_path)
print(f'Document saved to: {output_path}')
